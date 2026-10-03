from django.http import HttpResponse
from django.views.decorators.csrf import csrf_exempt

from pptx import Presentation
from pptx.util import Inches

from openpyxl import Workbook, load_workbook

from PIL import Image

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

import os
import tempfile
import zipfile
import subprocess

import pymupdf

from pypdf import PdfReader, PdfWriter

from pdf2docx import Converter

from docx import Document


# =========================================================
# HELPER FUNCTIONS
# =========================================================


def save_uploaded_file(uploaded_file, path):

    with open(path, "wb") as file:

        for chunk in uploaded_file.chunks():

            file.write(chunk)


def create_response(
    data,
    content_type,
    filename
):

    response = HttpResponse(
        data,
        content_type=content_type
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )

    return response


# =========================================================
# PDF → TEXT
# =========================================================


def convert_pdf_to_text(pdf):

    reader = PdfReader(pdf)

    all_text = ""

    for page in reader.pages:

        text = page.extract_text()

        if text:

            all_text += text + "\n"

    return create_response(
        all_text.encode("utf-8"),
        "text/plain",
        "converted.txt"
    )


# =========================================================
# PDF → WORD
# =========================================================


def convert_pdf_to_word(pdf):

    with tempfile.TemporaryDirectory() as temp_dir:

        pdf_path = os.path.join(
            temp_dir,
            "input.pdf"
        )

        word_path = os.path.join(
            temp_dir,
            "output.docx"
        )

        save_uploaded_file(
            pdf,
            pdf_path
        )

        converter = Converter(
            pdf_path
        )

        try:

            converter.convert(
                word_path
            )

        finally:

            converter.close()

        with open(
            word_path,
            "rb"
        ) as file:

            word_file = file.read()

        return create_response(
            word_file,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "converted.docx"
        )


# =========================================================
# PDF → EXCEL
# =========================================================


def convert_pdf_to_excel(pdf):

    with tempfile.TemporaryDirectory() as temp_dir:

        pdf_path = os.path.join(
            temp_dir,
            "input.pdf"
        )

        excel_path = os.path.join(
            temp_dir,
            "converted.xlsx"
        )

        save_uploaded_file(
            pdf,
            pdf_path
        )

        document = pymupdf.open(
            pdf_path
        )

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = "PDF Data"

        row_number = 1

        for page_number, page in enumerate(
            document,
            start=1
        ):

            worksheet.cell(
                row=row_number,
                column=1,
                value=f"Page {page_number}"
            )

            row_number += 1

            text = page.get_text(
                "text"
            )

            for line in text.splitlines():

                worksheet.cell(
                    row=row_number,
                    column=1,
                    value=line
                )

                row_number += 1

        document.close()

        workbook.save(
            excel_path
        )

        with open(
            excel_path,
            "rb"
        ) as file:

            excel_file = file.read()

        return create_response(
            excel_file,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "converted.xlsx"
        )


# =========================================================
# PDF → POWERPOINT
# =========================================================


def convert_pdf_to_powerpoint(pdf):

    with tempfile.TemporaryDirectory() as temp_dir:

        pdf_path = os.path.join(
            temp_dir,
            "input.pdf"
        )

        ppt_path = os.path.join(
            temp_dir,
            "output.pptx"
        )

        save_uploaded_file(
            pdf,
            pdf_path
        )

        document = pymupdf.open(
            pdf_path
        )

        presentation = Presentation()

        presentation.slide_width = Inches(
            13.333
        )

        presentation.slide_height = Inches(
            7.5
        )

        for page in document:

            pixmap = page.get_pixmap(
                matrix=pymupdf.Matrix(2, 2)
            )

            image_path = os.path.join(
                temp_dir,
                f"page-{page.number + 1}.png"
            )

            pixmap.save(
                image_path
            )

            slide = presentation.slides.add_slide(
                presentation.slide_layouts[6]
            )

            slide.shapes.add_picture(
                image_path,
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height
            )

        document.close()

        presentation.save(
            ppt_path
        )

        with open(
            ppt_path,
            "rb"
        ) as file:

            ppt_file = file.read()

        return create_response(
            ppt_file,
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "converted.pptx"
        )


# =========================================================
# PDF → JPG
# =========================================================


def convert_pdf_to_jpg(pdf):

    with tempfile.TemporaryDirectory() as temp_dir:

        pdf_path = os.path.join(
            temp_dir,
            "input.pdf"
        )

        zip_path = os.path.join(
            temp_dir,
            "images.zip"
        )

        save_uploaded_file(
            pdf,
            pdf_path
        )

        document = pymupdf.open(
            pdf_path
        )

        with zipfile.ZipFile(
            zip_path,
            "w",
            zipfile.ZIP_DEFLATED
        ) as zip_file:

            for page_number, page in enumerate(
                document
            ):

                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(2, 2)
                )

                image_path = os.path.join(
                    temp_dir,
                    f"page-{page_number + 1}.jpg"
                )

                pixmap.save(
                    image_path
                )

                zip_file.write(
                    image_path,
                    f"page-{page_number + 1}.jpg"
                )

        document.close()

        with open(
            zip_path,
            "rb"
        ) as file:

            zip_file_data = file.read()

        return create_response(
            zip_file_data,
            "application/zip",
            "pdf-images.zip"
        )


# =========================================================
# PDF → PNG
# =========================================================


def convert_pdf_to_png(pdf):

    with tempfile.TemporaryDirectory() as temp_dir:

        pdf_path = os.path.join(
            temp_dir,
            "input.pdf"
        )

        zip_path = os.path.join(
            temp_dir,
            "images.zip"
        )

        save_uploaded_file(
            pdf,
            pdf_path
        )

        document = pymupdf.open(
            pdf_path
        )

        with zipfile.ZipFile(
            zip_path,
            "w",
            zipfile.ZIP_DEFLATED
        ) as zip_file:

            for page_number, page in enumerate(
                document
            ):

                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(2, 2)
                )

                image_path = os.path.join(
                    temp_dir,
                    f"page-{page_number + 1}.png"
                )

                pixmap.save(
                    image_path
                )

                zip_file.write(
                    image_path,
                    f"page-{page_number + 1}.png"
                )

        document.close()

        with open(
            zip_path,
            "rb"
        ) as file:

            zip_file_data = file.read()

        return create_response(
            zip_file_data,
            "application/zip",
            "pdf-images-png.zip"
        )


# =========================================================
# JPG / PNG → PDF
# =========================================================


def convert_image_to_pdf(image_file):

    with tempfile.TemporaryDirectory() as temp_dir:

        image_path = os.path.join(
            temp_dir,
            image_file.name
        )

        pdf_path = os.path.join(
            temp_dir,
            "converted.pdf"
        )

        save_uploaded_file(
            image_file,
            image_path
        )

        image = Image.open(
            image_path
        )

        # Convert to RGB because
        # PDF does not support RGBA directly.

        if image.mode != "RGB":

            background = Image.new(
                "RGB",
                image.size,
                "white"
            )

            if "A" in image.getbands():

                background.paste(
                    image,
                    mask=image.getchannel("A")
                )

            else:

                background.paste(
                    image
                )

            image = background

        image.save(
            pdf_path,
            "PDF",
            resolution=100.0
        )

        with open(
            pdf_path,
            "rb"
        ) as file:

            pdf_file = file.read()

        return create_response(
            pdf_file,
            "application/pdf",
            "converted.pdf"
        )


# =========================================================
# TEXT → PDF
# =========================================================


def convert_text_to_pdf(text_file):

    with tempfile.TemporaryDirectory() as temp_dir:

        text_path = os.path.join(
            temp_dir,
            "input.txt"
        )

        pdf_path = os.path.join(
            temp_dir,
            "converted.pdf"
        )

        save_uploaded_file(
            text_file,
            text_path
        )

        try:

            with open(
                text_path,
                "r",
                encoding="utf-8",
                errors="replace"
            ) as file:

                lines = file.readlines()

        except Exception as error:

            return HttpResponse(
                f"Could not read text file: {error}",
                status=400
            )

        pdf = canvas.Canvas(
            pdf_path,
            pagesize=letter
        )

        width, height = letter

        x = 50

        y = height - 50

        line_height = 14

        for line in lines:

            line = line.rstrip("\n")

            # Basic line wrapping

            while len(line) > 100:

                part = line[:100]

                pdf.drawString(
                    x,
                    y,
                    part
                )

                y -= line_height

                line = line[100:]

                if y < 50:

                    pdf.showPage()

                    y = height - 50

            pdf.drawString(
                x,
                y,
                line
            )

            y -= line_height

            if y < 50:

                pdf.showPage()

                y = height - 50

        pdf.save()

        with open(
            pdf_path,
            "rb"
        ) as file:

            pdf_file = file.read()

        return create_response(
            pdf_file,
            "application/pdf",
            "converted.pdf"
        )


# =========================================================
# WORD → TEXT
# =========================================================


def convert_word_to_text(word_file):

    with tempfile.TemporaryDirectory() as temp_dir:

        word_path = os.path.join(
            temp_dir,
            "input.docx"
        )

        save_uploaded_file(
            word_file,
            word_path
        )

        document = Document(
            word_path
        )

        text = ""

        for paragraph in document.paragraphs:

            text += paragraph.text + "\n"

        return create_response(
            text.encode("utf-8"),
            "text/plain",
            "converted.txt"
        )


# =========================================================
# EXCEL → TEXT
# =========================================================


def convert_excel_to_text(excel_file):

    with tempfile.TemporaryDirectory() as temp_dir:

        excel_path = os.path.join(
            temp_dir,
            "input.xlsx"
        )

        save_uploaded_file(
            excel_file,
            excel_path
        )

        workbook = load_workbook(
            excel_path,
            data_only=True
        )

        output = []

        for worksheet in workbook.worksheets:

            output.append(
                f"--- {worksheet.title} ---"
            )

            for row in worksheet.iter_rows(
                values_only=True
            ):

                values = []

                for value in row:

                    if value is not None:

                        values.append(
                            str(value)
                        )

                    else:

                        values.append("")
                
                output.append(
                    "\t".join(values)
                )

        text = "\n".join(
            output
        )

        return create_response(
            text.encode("utf-8"),
            "text/plain",
            "converted.txt"
        )


# =========================================================
# OFFICE → PDF
#
# Requires LibreOffice installed on the system.
# =========================================================


def convert_office_to_pdf(uploaded_file):

    with tempfile.TemporaryDirectory() as temp_dir:

        input_path = os.path.join(
            temp_dir,
            uploaded_file.name
        )

        save_uploaded_file(
            uploaded_file,
            input_path
        )

        try:

            subprocess.run(
                [
                    "libreoffice",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    temp_dir,
                    input_path
                ],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE
            )

        except FileNotFoundError:

            return HttpResponse(
                "LibreOffice is not installed. "
                "Install LibreOffice to enable Office → PDF conversion.",
                status=500
            )

        except subprocess.CalledProcessError as error:

            return HttpResponse(
                "Office to PDF conversion failed.",
                status=500
            )

        base_name = os.path.splitext(
            os.path.basename(uploaded_file.name)
        )[0]

        pdf_path = os.path.join(
            temp_dir,
            base_name + ".pdf"
        )

        if not os.path.exists(pdf_path):

            return HttpResponse(
                "LibreOffice did not create a PDF.",
                status=500
            )

        with open(
            pdf_path,
            "rb"
        ) as file:

            pdf_file = file.read()

        return create_response(
            pdf_file,
            "application/pdf",
            "converted.pdf"
        )


# =========================================================
# LEGACY PDF → TEXT ENDPOINT
# =========================================================


@csrf_exempt
def pdf_to_text(request):

    if request.method != "POST":

        return HttpResponse(
            "POST request required",
            status=405
        )

    if "pdf" not in request.FILES:

        return HttpResponse(
            "No PDF file uploaded",
            status=400
        )

    return convert_pdf_to_text(
        request.FILES["pdf"]
    )


# =========================================================
# LEGACY PDF → WORD ENDPOINT
# =========================================================


@csrf_exempt
def pdf_to_word(request):

    if request.method != "POST":

        return HttpResponse(
            "POST request required",
            status=405
        )

    if "pdf" not in request.FILES:

        return HttpResponse(
            "No PDF file uploaded",
            status=400
        )

    return convert_pdf_to_word(
        request.FILES["pdf"]
    )


# =========================================================
# LEGACY PDF → IMAGE ENDPOINT
# =========================================================


@csrf_exempt
def pdf_to_image(request):

    if request.method != "POST":

        return HttpResponse(
            "POST request required",
            status=405
        )

    if "pdf" not in request.FILES:

        return HttpResponse(
            "No PDF file uploaded",
            status=400
        )

    return convert_pdf_to_jpg(
        request.FILES["pdf"]
    )


# =========================================================
# SPLIT PDF
# =========================================================


@csrf_exempt
def pdf_to_split(request):

    if request.method != "POST":

        return HttpResponse(
            "POST request required",
            status=405
        )

    if "pdf" not in request.FILES:

        return HttpResponse(
            "No PDF file uploaded",
            status=400
        )

    pdf = request.FILES["pdf"]

    try:

        start_page = int(
            request.POST["start_page"]
        )

        end_page = int(
            request.POST["end_page"]
        )

    except (
        KeyError,
        ValueError
    ):

        return HttpResponse(
            "Valid start_page and end_page are required.",
            status=400
        )

    reader = PdfReader(
        pdf
    )

    total_pages = len(
        reader.pages
    )

    if start_page < 1:

        return HttpResponse(
            "Start page must be at least 1.",
            status=400
        )

    if end_page > total_pages:

        return HttpResponse(
            f"Invalid page range. PDF has {total_pages} pages.",
            status=400
        )

    if start_page > end_page:

        return HttpResponse(
            "Start page cannot be greater than end page.",
            status=400
        )

    writer = PdfWriter()

    for page_number in range(
        start_page - 1,
        end_page
    ):

        writer.add_page(
            reader.pages[page_number]
        )

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="split.pdf"'
    )

    writer.write(
        response
    )

    return response


# =========================================================
# MERGE PDF
# =========================================================


@csrf_exempt
def pdf_to_merge(request):

    if request.method != "POST":

        return HttpResponse(
            "POST request required",
            status=405
        )

    pdf_files = request.FILES.getlist(
        "pdfs"
    )

    if len(pdf_files) < 2:

        return HttpResponse(
            "Please select at least two PDF files.",
            status=400
        )

    writer = PdfWriter()

    for pdf in pdf_files:

        reader = PdfReader(
            pdf
        )

        for page in reader.pages:

            writer.add_page(
                page
            )

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="merged.pdf"'
    )

    writer.write(
        response
    )

    return response


# =========================================================
# PDF → POWERPOINT ENDPOINT
# =========================================================


@csrf_exempt
def pdf_to_ppt(request):

    if request.method != "POST":

        return HttpResponse(
            "POST request required",
            status=405
        )

    if "pdf" not in request.FILES:

        return HttpResponse(
            "No PDF file uploaded",
            status=400
        )

    return convert_pdf_to_powerpoint(
        request.FILES["pdf"]
    )


# =========================================================
# PDF → EXCEL ENDPOINT
# =========================================================


@csrf_exempt
def pdf_to_excel(request):

    if request.method != "POST":

        return HttpResponse(
            "POST request required",
            status=405
        )

    if "pdf" not in request.FILES:

        return HttpResponse(
            "No PDF file uploaded",
            status=400
        )

    return convert_pdf_to_excel(
        request.FILES["pdf"]
    )


# =========================================================
# UNIVERSAL CONVERTER
# =========================================================


@csrf_exempt
def convert(request):

    if request.method != "POST":

        return HttpResponse(
            "POST request required",
            status=405
        )


    # =====================================================
    # FILE
    # =====================================================

    if "file" not in request.FILES:

        return HttpResponse(
            "No file uploaded.",
            status=400
        )

    file = request.FILES["file"]


    # =====================================================
    # FORMATS
    # =====================================================

    from_format = (
        request.POST
        .get("from_format", "")
        .strip()
    )

    to_format = (
        request.POST
        .get("to_format", "")
        .strip()
    )


    print(
        "From:",
        from_format
    )

    print(
        "To:",
        to_format
    )

    print(
        "File:",
        file.name
    )


    if not from_format:

        return HttpResponse(
            "Source format is required.",
            status=400
        )


    if not to_format:

        return HttpResponse(
            "Target format is required.",
            status=400
        )


    # =====================================================
    # PDF → TEXT
    # =====================================================

    if (
        from_format == "PDF"
        and
        to_format == "Text"
    ):

        return convert_pdf_to_text(
            file
        )


    # =====================================================
    # PDF → WORD
    # =====================================================

    if (
        from_format == "PDF"
        and
        to_format == "Word"
    ):

        return convert_pdf_to_word(
            file
        )


    # =====================================================
    # PDF → EXCEL
    # =====================================================

    if (
        from_format == "PDF"
        and
        to_format == "Excel"
    ):

        return convert_pdf_to_excel(
            file
        )


    # =====================================================
    # PDF → POWERPOINT
    # =====================================================

    if (
        from_format == "PDF"
        and
        to_format == "PowerPoint"
    ):

        return convert_pdf_to_powerpoint(
            file
        )


    # =====================================================
    # PDF → JPG
    # =====================================================

    if (
        from_format == "PDF"
        and
        to_format == "JPG"
    ):

        return convert_pdf_to_jpg(
            file
        )


    # =====================================================
    # PDF → PNG
    # =====================================================

    if (
        from_format == "PDF"
        and
        to_format == "PNG"
    ):

        return convert_pdf_to_png(
            file
        )


    # =====================================================
    # JPG → PDF
    # =====================================================

    if (
        from_format == "JPG"
        and
        to_format == "PDF"
    ):

        return convert_image_to_pdf(
            file
        )


    # =====================================================
    # PNG → PDF
    # =====================================================

    if (
        from_format == "PNG"
        and
        to_format == "PDF"
    ):

        return convert_image_to_pdf(
            file
        )


    # =====================================================
    # JPEG → PDF
    # =====================================================

    if (
        from_format == "JPEG"
        and
        to_format == "PDF"
    ):

        return convert_image_to_pdf(
            file
        )


    # =====================================================
    # TEXT → PDF
    # =====================================================

    if (
        from_format == "Text"
        and
        to_format == "PDF"
    ):

        return convert_text_to_pdf(
            file
        )


    # =====================================================
    # WORD → TEXT
    # =====================================================

    if (
        from_format == "Word"
        and
        to_format == "Text"
    ):

        return convert_word_to_text(
            file
        )


    # =====================================================
    # EXCEL → TEXT
    # =====================================================

    if (
        from_format == "Excel"
        and
        to_format == "Text"
    ):

        return convert_excel_to_text(
            file
        )


    # =====================================================
    # WORD → PDF
    # =====================================================

    if (
        from_format == "Word"
        and
        to_format == "PDF"
    ):

        return convert_office_to_pdf(
            file
        )


    # =====================================================
    # EXCEL → PDF
    # =====================================================

    if (
        from_format == "Excel"
        and
        to_format == "PDF"
    ):

        return convert_office_to_pdf(
            file
        )


    # =====================================================
    # POWERPOINT → PDF
    # =====================================================

    if (
        from_format == "PowerPoint"
        and
        to_format == "PDF"
    ):

        return convert_office_to_pdf(
            file
        )


    # =====================================================
    # UNSUPPORTED
    # =====================================================

    return HttpResponse(
        f"{from_format} → {to_format} is not available yet.",
        status=400
    )